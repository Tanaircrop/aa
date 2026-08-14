"""Kiểm tra API end-to-end trên một DB tạm: import, nhập liệu, IRR, export."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    """App chạy trên DB tạm, seed bằng workbook demo sinh tại chỗ."""
    db_path = tmp_path_factory.mktemp("db") / "test.db"
    import os

    os.environ["TIKTOK_CODING_DB"] = str(db_path)
    for module in [m for m in list(sys.modules) if m.startswith("backend")]:
        del sys.modules[module]

    from fastapi.testclient import TestClient

    from backend.app import app
    from backend.database import engine, init_db
    from backend.services.importer import import_workbook
    from seed.make_demo_xlsx import build
    from sqlmodel import Session

    xlsx = tmp_path_factory.mktemp("seed") / "demo.xlsx"
    build(24, 8).save(xlsx)

    init_db()
    with Session(engine) as session:
        import_workbook(session, xlsx)

    with TestClient(app) as test_client:
        test_client.xlsx_path = xlsx
        yield test_client


def first_uncoded(client) -> str:
    rows = client.get("/api/videos", params={"status": "not_started"}).json()
    return rows["videos"][0]["video_id"]


COMPLETE = {
    "coder_id": "C1", "account_type": 1, "x1_commercial": 1, "x2_4_hashtag": 1,
    "x4a_video_sec": 60, "x4b_commercial_sec": 30,
    "y1_information": 1, "y2_experience": 0, "y3_aesthetic": 0,
    "y4_price_promo": 0, "y5_social_proof": 0, "y6_problem_solution": 0,
    "y7_dominant_frame": 1, "y8_cta": 1, "y9_appeal": 4, "w1_ai_disclosure": 0,
}


# ------------------------------------------------------------------ Schema

def test_schema_exposes_all_47_columns(client):
    data = client.get("/api/meta/schema").json()
    assert len(data["columns"]) == 47
    assert data["columns"][0]["key"] == "stt"
    assert data["columns"][-1]["key"] == "boundary_notes"
    assert data["tab_order"]
    assert data["quick_guide"]["y7_dominant_frame"]["measures"]


def test_import_is_idempotent(client):
    from sqlmodel import Session, select

    from backend.database import engine
    from backend.models import CodingEntry, Video
    from backend.services.importer import import_workbook

    with Session(engine) as session:
        before = len(session.exec(select(Video)).all())
        import_workbook(session, client.xlsx_path)
        import_workbook(session, client.xlsx_path)
        after = len(session.exec(select(Video)).all())
        entries = session.exec(select(CodingEntry)).all()
    assert before == after
    assert len({e.video_id for e in entries}) == len(entries)


# ------------------------------------------------------------- Nhập liệu

def test_entry_autopulls_metadata(client):
    video_id = first_uncoded(client)
    payload = client.get(f"/api/coding/{video_id}", params={"coder": "C1"}).json()
    values = payload["values"]
    assert values["video_id"] == video_id
    assert values["url"]
    assert values["account_handle"]
    assert values["e4_views"] is not None
    assert values["coder_id"] == "C1"
    assert payload["status"] == "not_started"


def test_patch_autosaves_and_recomputes(client):
    video_id = first_uncoded(client)
    payload = client.patch(f"/api/coding/{video_id}",
                           json={"fields": COMPLETE}).json()
    assert payload["derived"]["values"]["x4c_commercial_share"] == 50.0
    assert payload["derived"]["values"]["qc_status"] == "OK"
    assert payload["status"] == "ok"
    assert payload["completed"] is True

    again = client.get(f"/api/coding/{video_id}").json()
    assert again["values"]["y9_appeal"] == 4


def test_partial_entry_reports_missing_fields(client):
    video_id = first_uncoded(client)
    payload = client.patch(f"/api/coding/{video_id}",
                           json={"fields": {"x1_commercial": 1}}).json()
    assert payload["status"] == "in_progress"
    assert payload["validation"]["can_complete"] is False
    assert "y1_information" in payload["validation"]["missing_required"]


def test_hard_warning_does_not_block_saving(client):
    video_id = first_uncoded(client)
    fields = dict(COMPLETE, x4a_video_sec=10, x4b_commercial_sec=50)
    payload = client.patch(f"/api/coding/{video_id}", json={"fields": fields}).json()
    assert payload["validation"]["has_hard"] is True
    assert payload["values"]["x4b_commercial_sec"] == 50   # vẫn lưu
    assert payload["derived"]["values"]["qc_status"] == "CHECK"


def test_override_derived_field(client):
    video_id = first_uncoded(client)
    client.patch(f"/api/coding/{video_id}", json={"fields": COMPLETE})
    payload = client.patch(f"/api/coding/{video_id}",
                           json={"overrides": {"x4_intensity_band": 0}}).json()
    assert payload["derived"]["values"]["x4_intensity_band"] == 0
    assert payload["derived"]["meta"]["x4_intensity_band"]["overridden"] is True

    cleared = client.patch(f"/api/coding/{video_id}",
                           json={"overrides": {"x4_intensity_band": None}}).json()
    assert cleared["derived"]["meta"]["x4_intensity_band"]["overridden"] is False


def test_unknown_field_is_rejected_not_crashed(client):
    video_id = first_uncoded(client)
    payload = client.patch(f"/api/coding/{video_id}",
                           json={"fields": {"khong_ton_tai": 1}}).json()
    assert "khong_ton_tai" in payload["rejected_fields"]


def test_neighbors_follow_active_filter(client):
    rows = client.get("/api/videos", params={"status": "not_started"}).json()["videos"]
    assert len(rows) >= 2
    nav = client.get(f"/api/videos/{rows[0]['video_id']}/neighbors",
                     params={"status": "not_started"}).json()
    assert nav["next"] == rows[1]["video_id"]
    assert nav["prev"] is None
    assert nav["index"] == 1


def test_flag_for_review(client):
    video_id = first_uncoded(client)
    payload = client.patch(f"/api/coding/{video_id}",
                           json={"fields": {}, "flagged_for_review": True}).json()
    assert payload["flagged_for_review"] is True
    flagged = client.get("/api/videos", params={"status": "flagged"}).json()
    assert video_id in [v["video_id"] for v in flagged["videos"]]


# ------------------------------------------------------------- Dashboard

def test_dashboard_summary_and_charts(client):
    summary = client.get("/api/dashboard/summary").json()
    assert summary["target"] >= summary["queue_total"]
    assert summary["coded"] >= 1
    assert summary["qc"]["ok"] + summary["qc"]["check"] == summary["coded"]

    charts = client.get("/api/dashboard/charts").json()
    assert charts["y7"]["categories"]
    assert len(charts["y7"]["values"]) == len(charts["y7"]["labels"])


def test_untouched_videos_are_not_in_review_queue(client):
    review = client.get("/api/dashboard/needs-review").json()
    started = {v["video_id"] for v in client.get("/api/videos").json()["videos"]
               if v["status"] != "not_started"}
    assert {r["video_id"] for r in review} <= started


def test_grid_inline_edit(client):
    video_id = first_uncoded(client)
    client.patch(f"/api/coding/{video_id}", json={"fields": COMPLETE})
    res = client.patch("/api/dashboard/grid",
                       json={"video_id": video_id, "field": "y8_cta", "value": 2})
    assert res.json()["values"]["y8_cta"] == 2


# ------------------------------------------------------------------- IRR

def test_irr_roles_are_independent(client):
    pilot = client.get("/api/irr/videos", params={"role": "C1"}).json()
    video_id = pilot["videos"][0]["video_id"]

    c1 = client.patch(f"/api/irr/C1/{video_id}", json={"fields": COMPLETE}).json()
    assert c1["completed"] is True

    # C2 mở cùng video: không thấy giá trị nào của C1.
    c2 = client.get(f"/api/irr/C2/{video_id}").json()
    assert c2["values"]["y9_appeal"] is None
    assert c2["values"]["x1_commercial"] is None
    assert c2["values"]["url"]  # metadata thì vẫn auto-pull


def test_compare_blocked_until_both_coded(client):
    pilot = client.get("/api/irr/videos", params={"role": "C1"}).json()
    video_id = pilot["videos"][1]["video_id"]
    client.patch(f"/api/irr/C1/{video_id}", json={"fields": COMPLETE})

    blocked = client.get(f"/api/irr/compare/{video_id}")
    assert blocked.status_code == 409

    client.patch(f"/api/irr/C2/{video_id}",
                 json={"fields": dict(COMPLETE, y8_cta=2, y9_appeal=1)})
    data = client.get(f"/api/irr/compare/{video_id}").json()
    assert data["agreement"]["total"] > 0
    diffs = {r["field"] for r in data["rows"] if not r["agree"]}
    assert diffs == {"y8_cta", "y9_appeal"}


def test_irr_scale_types_match_codebook(client):
    pilot = client.get("/api/irr/videos", params={"role": "C1"}).json()
    video_id = pilot["videos"][1]["video_id"]
    rows = {r["field"]: r["scale"] for r in
            client.get(f"/api/irr/compare/{video_id}").json()["rows"]}
    assert rows["x4a_video_sec"] == "interval"
    assert rows["x4b_commercial_sec"] == "interval"
    assert rows["y8_cta"] == "ordinal"
    assert rows["y7_dominant_frame"] == "nominal"


def test_irr_matrix_export_is_wide_with_two_columns_per_variable(client):
    text = client.get("/api/irr/export/matrix.csv").text
    header = text.splitlines()[0].split(",")
    assert "y8_cta_C1" in header and "y8_cta_C2" in header
    assert text.splitlines()[1].split(",")[1] == "scale_type"


def test_irr_long_export(client):
    lines = client.get("/api/irr/export/long.csv").text.splitlines()
    assert lines[0] == "unit,coder,variable,scale_type,value"
    assert len(lines) > 1


def test_invalid_role_rejected(client):
    assert client.get("/api/irr/C9/abc").status_code == 400


# ---------------------------------------------------------------- Export

def test_export_xlsx_has_47_columns_in_order(client):
    import io

    from openpyxl import load_workbook

    from backend.columns import EXCEL_HEADERS

    res = client.get("/api/io/export/coding.xlsx")
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(res.content))
    ws = wb["3_Coding_Sheet"]
    headers = [c.value for c in ws[1]]
    assert headers == EXCEL_HEADERS
    assert ws.max_row > 1


def test_export_full_includes_irr_sheet(client):
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(client.get("/api/io/export/full.xlsx").content))
    assert "3_Coding_Sheet" in wb.sheetnames
    assert "5_IRR_Pilot" in wb.sheetnames


def test_export_csv(client):
    text = client.get("/api/io/export/coding.csv").content.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("STT,")


# ------------------------------------------------------------------ Sync

def test_sync_reports_missing_configuration_clearly(client):
    res = client.get("/api/sync/pull/preview")
    assert res.status_code == 400
    assert "gspread" in res.json()["detail"] or "cấu hình" in res.json()["detail"]
