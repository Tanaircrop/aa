"""Đọc/ghi file `.xlsx` (mục 10 của SPEC).

Import chấp nhận header hơi lệch so với bản gốc (khác dấu cách, thiếu tiền tố,
khác hoa/thường) nhờ chuẩn hoá + bảng alias, vì file thật thường bị chỉnh tay.
Export luôn ghi đúng thứ tự và đúng tên 47 cột của `3_Coding_Sheet`.
"""

from __future__ import annotations

import io
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..columns import COLUMNS, COLUMNS_BY_KEY, EXCEL_HEADERS

SHEET_ACCOUNTS = "1_Account_List"
SHEET_QUEUE = "2_Video_Queue"
SHEET_CODING = "3_Coding_Sheet"
SHEET_IRR = "5_IRR_Pilot"
SHEET_GUIDE = "Y_Frame_Quick_Guide"


def normalize(text: Any) -> str:
    """Chuẩn hoá header để so khớp: bỏ dấu, bỏ ký tự không phải chữ/số.

    Lưu ý: `đ`/`Đ` là ký tự độc lập (U+0111/U+0110), NFKD không tách được thành
    `d` + dấu, nên phải thay tay — nếu không, header kiểu "Đo cái gì?" sẽ mất chữ
    đầu và không khớp với gì cả.
    """
    if text is None:
        return ""
    s = str(text).replace("đ", "d").replace("Đ", "D")
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]", "", s.lower())


# Header phụ thường gặp trong file thật -> key nội bộ
HEADER_ALIASES: dict[str, str] = {
    "videoid": "video_id",
    "v01videoid": "video_id",
    "v1videoid": "video_id",
    "url": "url",
    "videourl": "url",
    "link": "url",
    "uploaddate": "upload_date",
    "accounthandle": "account_handle",
    "handle": "account_handle",
    "account": "account_handle",
    "coderid": "coder_id",
    "coder": "coder_id",
    "codingdate": "coding_date",
    "snapshottime": "snapshot_time",
    "accounttype": "account_type",
    "followercount": "follower_count",
    "follower": "follower_count",
    "followers": "follower_count",
    "likes": "e1_likes",
    "comments": "e2_comments",
    "shares": "e3_shares",
    "views": "e4_views",
    "saves": "e5_saves",
    "stt": "stt",
    "qcstatus": "qc_status",
}

_COLUMN_LOOKUP: dict[str, str] = {}
for _c in COLUMNS:
    _COLUMN_LOOKUP[normalize(_c.excel)] = _c.key
    _COLUMN_LOOKUP[normalize(_c.key)] = _c.key
    if _c.alias:
        _COLUMN_LOOKUP.setdefault(normalize(_c.alias), _c.key)
_COLUMN_LOOKUP.update(HEADER_ALIASES)


def resolve_header(header: Any) -> str | None:
    return _COLUMN_LOOKUP.get(normalize(header))


#: Số cột phải khớp tối thiểu thì một dòng mới được coi là header. Dưới ngưỡng
#: này (ví dụ sheet quick guide, không có cột nào trùng tên) thì dùng dòng đầu.
MIN_HEADER_HITS = 3


def find_header_row(rows: list[tuple[Any, ...]], max_scan: int = 10) -> int:
    """Tìm dòng header: dòng khớp nhiều tên cột đã biết nhất trong N dòng đầu.

    Nếu không dòng nào đạt `MIN_HEADER_HITS`, trả về 0 — tránh nhầm một dòng dữ
    liệu (ví dụ dòng bắt đầu bằng "X1") thành header.
    """
    best_idx, best_hits = 0, 0
    for idx, row in enumerate(rows[:max_scan]):
        hits = sum(1 for cell in row if resolve_header(cell))
        if hits > best_hits:
            best_idx, best_hits = idx, hits
    return best_idx if best_hits >= MIN_HEADER_HITS else 0


def read_sheet(path: str | Path, sheet_name: str,
               header_row: int | None = None) -> list[dict[str, Any]]:
    """Đọc một sheet -> list dict {key nội bộ: value}, kèm `_raw` cho cột lạ.

    `header_row` (0-based) để ép dòng header khi sheet không có cột nào trùng
    tên với 47 cột chuẩn.
    """
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        actual = _match_sheet(wb.sheetnames, sheet_name)
        if actual is None:
            return []
        ws = wb[actual]
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    if not rows:
        return []

    header_idx = header_row if header_row is not None else find_header_row(rows)
    header_idx = max(0, min(header_idx, len(rows) - 1))
    headers = rows[header_idx]
    out: list[dict[str, Any]] = []
    for row in rows[header_idx + 1:]:
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        record: dict[str, Any] = {"_raw": {}}
        for header, value in zip(headers, row):
            if header is None:
                continue
            key = resolve_header(header)
            if key:
                record[key] = _clean(value)
            else:
                record["_raw"][str(header).strip()] = _clean(value)
        out.append(record)
    return out


def _match_sheet(names: Iterable[str], wanted: str) -> str | None:
    target = normalize(wanted)
    names = list(names)
    for name in names:
        if normalize(name) == target:
            return name
    # Khớp lỏng: "3_Coding_Sheet" vs "Coding Sheet"
    core = re.sub(r"^\d+", "", target)
    for name in names:
        n = re.sub(r"^\d+", "", normalize(name))
        if core and (core in n or n in core):
            return name
    return None


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if isinstance(value, datetime):
        return value
    return value


def coerce(key: str, value: Any) -> Any:
    """Ép kiểu theo khai báo cột; giá trị không hợp lệ -> None (không crash)."""
    col = COLUMNS_BY_KEY.get(key)
    if col is None or value is None or value == "":
        return None
    dtype = col.dtype
    try:
        if dtype == "int":
            if isinstance(value, str):
                value = value.replace(",", "").replace(".", "").strip()
            return int(float(value))
        if dtype == "float":
            if isinstance(value, str):
                value = value.replace(",", "").strip()
            return float(value)
        if dtype in ("bool01",):
            return 1 if str(value).strip() in ("1", "True", "true", "x", "X") else 0
        if dtype == "enum":
            if col.options and isinstance(col.options[0], int):
                return int(float(value))
            return str(value).strip()
        if dtype == "date":
            return _to_date(value)
        if dtype == "datetime":
            return _to_datetime(value)
        return str(value).strip()
    except (TypeError, ValueError):
        return None


_DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%Y/%m/%d")


def _to_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def _to_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M",
                "%d/%m/%Y %H:%M"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    d = _to_date(value)
    return datetime(d.year, d.month, d.day) if d else None


# --------------------------------------------------------------------------
# Export
# --------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F3B57")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
CHECK_FILL = PatternFill("solid", fgColor="FFF2CC")


def _cell_value(row: dict[str, Any], key: str) -> Any:
    value = row.get(key)
    if isinstance(value, (datetime, date)):
        return value
    return value


def write_coding_sheet(ws, rows: list[dict[str, Any]]) -> None:
    ws.append(EXCEL_HEADERS)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ws.freeze_panes = "C2"

    qc_idx = EXCEL_HEADERS.index("QC_status") + 1
    for row in rows:
        ws.append([_cell_value(row, c.key) for c in COLUMNS])
        if str(row.get("qc_status") or "").upper() == "CHECK":
            ws.cell(row=ws.max_row, column=qc_idx).fill = CHECK_FILL

    for idx, col in enumerate(COLUMNS, start=1):
        width = 12
        if col.dtype in ("text",):
            width = 28 if "notes" in col.key else 22
        if col.key in ("url",):
            width = 42
        ws.column_dimensions[get_column_letter(idx)].width = width


def export_workbook(coding_rows: list[dict[str, Any]],
                    irr_rows: list[dict[str, Any]] | None = None) -> bytes:
    """Xuất workbook: sheet `3_Coding_Sheet` (+ `5_IRR_Pilot` nếu có)."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_CODING
    write_coding_sheet(ws, coding_rows)

    if irr_rows:
        ws2 = wb.create_sheet(SHEET_IRR)
        headers = ["Video_ID", "Role"] + EXCEL_HEADERS[1:] + ["Ghi_chu_khac_biet"]
        ws2.append(headers)
        for cell in ws2[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws2.freeze_panes = "C2"
        for row in irr_rows:
            values = [row.get("video_id"), row.get("role")]
            values += [_cell_value(row, c.key) for c in COLUMNS[1:]]
            values.append(row.get("diff_notes"))
            ws2.append(values)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def rows_to_csv(rows: list[dict[str, Any]]) -> str:
    import csv

    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(EXCEL_HEADERS)
    for row in rows:
        out = []
        for col in COLUMNS:
            value = row.get(col.key)
            if isinstance(value, (datetime, date)):
                value = value.isoformat()
            out.append("" if value is None else value)
        writer.writerow(out)
    return buffer.getvalue()
